from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter

INK="1F2937"; GREEN="1F5E3A"; AMBER="F59E0B"; CREAM="FFF7ED"; PALE="ECFDF5"; LIGHT="F8FAFC"; RED="FEE2E2"; AMB="FEF3C7"; GRN="DCFCE7"; GREY="9CA3AF"
thin=Side(style="thin",color="E5E7EB"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
def fill(c): return PatternFill("solid",fgColor=c)
def hdr(ws,row,cols,bg=GREEN,fg="FFFFFF"):
    for i,t in enumerate(cols,1):
        c=ws.cell(row=row,column=i,value=t); c.font=Font(bold=True,color=fg,size=11); c.fill=fill(bg); c.alignment=Alignment(vertical="center",wrap_text=True); c.border=border
    ws.row_dimensions[row].height=30

wb=Workbook()

# ---------- GUIDE ----------
g=wb.active; g.title="📖 Guide"
g.sheet_properties.tabColor=INK
g.column_dimensions["A"].width=4; g.column_dimensions["B"].width=34; g.column_dimensions["C"].width=90
g["B2"]="Idées de niches — carnet de recherche"; g["B2"].font=Font(bold=True,size=20,color=GREEN)
g["B3"]="Une feuille par niche. Dupliquer l'onglet « 🧩 MODÈLE », le renommer, puis l'ajouter à l'index."; g["B3"].font=Font(italic=True,color=GREY)
rows=[
("Comment ça marche",""),
("1. Dupliquer","Clic droit sur l'onglet « 🧩 MODÈLE » → Dupliquer → renommer avec le nom de la niche (ex. « Abri chat extérieur »)."),
("2. En-tête de niche","Remplir le bloc du haut : nom, mode (PRODUIT PUR ou UNIVERS), marché, statut, date. Le mode pilote les seuils automatiquement."),
("3. Arborescence","Une ligne par mot-clé. Colonne A = niveau : Mot-clé principal → Collection → Produit. Les couleurs se mettent toutes seules."),
("4. Volumes","Colonne D = volume mensuel DataForSEO (jamais SEMrush). Le seuil (E) et le verdict (F) se calculent seuls : PASS ≥ seuil, REVIEW dans la bande −20 %, STOP en dessous."),
("5. Produits","Pour chaque ligne Produit : lien AliExpress (G), prix fournisseur (H), prix cible (I), marge indicative (J, calculée). Un lien suffit, le sourcing complet vient après le PASS."),
("6. Index","Dans « 🗂 Index », taper le nom exact de l'onglet en colonne A : mode, statut, mot-clé principal, volume, compteurs et verdict remontent automatiquement."),
("",""),
("Seuils du parc (DataForSEO, France ; appliqués tels quels au UK)",""),
("PRODUIT PUR","Mot-clé principal ≥ 12 500 recherches/mois (cluster adressable, MAX du groupe, jamais la somme)."),
("UNIVERS","Volume consolidé des familles ≥ 37 500/mois (confort 50 000). Une tête seule ne mesure pas un univers."),
("Bande REVIEW","Entre 80 % et 100 % du seuil → REVIEW, à remonter, jamais tranché en silence."),
("Règle d'or","On ne somme jamais des volumes bruts Google : les variantes proches sont déjà agrégées. On prend le MAX par groupe."),
("",""),
("Niveaux et couleurs",""),
("Mot-clé principal","Ligne vert foncé — la tête de la niche. Une seule par feuille (deux si deux têtes candidates)."),
("Collection","Ligne vert pâle — un rayon de la boutique (mode UNIVERS) ou une famille du produit (mode PRODUIT PUR)."),
("Produit","Ligne blanche — un produit précis, avec son lien AliExpress."),
("",""),
("Statuts","Idée → Mesure en cours → PASS → REVIEW → STOP → Sourcing → GO_FINAL / NO_GO_FINAL (les GO restent une décision de Hakim)."),
]
r=5
for a,b in rows:
    g.cell(row=r,column=2,value=a); g.cell(row=r,column=3,value=b)
    if b=="" and a: g.cell(row=r,column=2).font=Font(bold=True,size=13,color=GREEN)
    else: g.cell(row=r,column=2).font=Font(bold=True,color=INK)
    g.cell(row=r,column=3).alignment=Alignment(wrap_text=True,vertical="top"); g.cell(row=r,column=2).alignment=Alignment(vertical="top")
    r+=1
# legend swatches
g["B28"]="Légende verdict"; g["B28"].font=Font(bold=True,size=13,color=GREEN)
for i,(t,c) in enumerate([("PASS",GRN),("REVIEW",AMB),("STOP",RED)]):
    cell=g.cell(row=29+i,column=2,value=t); cell.fill=fill(c); cell.font=Font(bold=True); cell.border=border

# ---------- INDEX ----------
ix=wb.create_sheet("🗂 Index"); ix.sheet_properties.tabColor=AMBER
ix["A1"]="Index des niches"; ix["A1"].font=Font(bold=True,size=18,color=GREEN)
ix["A2"]="Colonne A = nom exact de l'onglet. Le reste remonte tout seul."; ix["A2"].font=Font(italic=True,color=GREY)
cols=["Onglet (nom exact)","Mode","Marché","Statut","Mot-clé principal","Volume principal","Seuil","Verdict","Collections","Produits","Produits avec lien","Dernière mise à jour","Note"]
hdr(ix,4,cols)
widths=[30,14,10,18,30,14,10,11,12,10,14,16,40]
for i,w in enumerate(widths,1): ix.column_dimensions[get_column_letter(i)].width=w
for r in range(5,31):
    A=f"$A{r}"
    def ind(ref): return f'=IFERROR(IF({A}="","",INDIRECT("\'"&{A}&"\'!{ref}")),"")'
    ix.cell(row=r,column=2,value=ind("C3"))
    ix.cell(row=r,column=3,value=ind("C4"))
    ix.cell(row=r,column=4,value=ind("F3"))
    ix.cell(row=r,column=5,value=f'=IFERROR(IF({A}="","",INDEX(INDIRECT("\'"&{A}&"\'!C8:C80"),MATCH("Mot-clé principal",INDIRECT("\'"&{A}&"\'!A8:A80"),0))),"")')
    ix.cell(row=r,column=6,value=f'=IFERROR(IF({A}="","",INDEX(INDIRECT("\'"&{A}&"\'!D8:D80"),MATCH("Mot-clé principal",INDIRECT("\'"&{A}&"\'!A8:A80"),0))),"")')
    ix.cell(row=r,column=7,value=f'=IFERROR(IF({A}="","",INDEX(INDIRECT("\'"&{A}&"\'!E8:E80"),MATCH("Mot-clé principal",INDIRECT("\'"&{A}&"\'!A8:A80"),0))),"")')
    ix.cell(row=r,column=8,value=f'=IFERROR(IF({A}="","",INDEX(INDIRECT("\'"&{A}&"\'!F8:F80"),MATCH("Mot-clé principal",INDIRECT("\'"&{A}&"\'!A8:A80"),0))),"")')
    ix.cell(row=r,column=9,value=f'=IFERROR(IF({A}="","",COUNTIF(INDIRECT("\'"&{A}&"\'!A8:A80"),"Collection")),"")')
    ix.cell(row=r,column=10,value=f'=IFERROR(IF({A}="","",COUNTIF(INDIRECT("\'"&{A}&"\'!A8:A80"),"Produit")),"")')
    ix.cell(row=r,column=11,value=f'=IFERROR(IF({A}="","",COUNTIFS(INDIRECT("\'"&{A}&"\'!A8:A80"),"Produit",INDIRECT("\'"&{A}&"\'!G8:G80"),"<>")),"")')
    ix.cell(row=r,column=12,value=ind("F4"))
    for c in range(1,14):
        cell=ix.cell(row=r,column=c); cell.border=border
        if r%2==0: cell.fill=fill(LIGHT)
    ix.cell(row=r,column=6).number_format='# ##0'; ix.cell(row=r,column=7).number_format='# ##0'
ix.freeze_panes="B5"
ix.conditional_formatting.add("H5:H30",CellIsRule(operator="equal",formula=['"PASS"'],fill=fill(GRN),font=Font(bold=True,color="166534")))
ix.conditional_formatting.add("H5:H30",CellIsRule(operator="equal",formula=['"REVIEW"'],fill=fill(AMB),font=Font(bold=True,color="92400E")))
ix.conditional_formatting.add("H5:H30",CellIsRule(operator="equal",formula=['"STOP"'],fill=fill(RED),font=Font(bold=True,color="991B1B")))
dv=DataValidation(type="list",formula1='"Idée,Mesure en cours,PASS,REVIEW,STOP,Sourcing,GO_FINAL,WATCH_FINAL,NO_GO_FINAL"',allow_blank=True)

# ---------- NICHE SHEET BUILDER ----------
def niche_sheet(title,tab,example=None):
    ws=wb.create_sheet(title); ws.sheet_properties.tabColor=tab
    for col,w in zip("ABCDEFGHIJK",[18,30,34,12,10,10,44,12,12,12,40]): ws.column_dimensions[col].width=w
    ws["A1"]="NICHE"; ws["A1"].font=Font(bold=True,size=9,color=GREY)
    ws["B1"]=None
    ws.merge_cells("B2:D2"); ws["B2"]=example["name"] if example else "Nom de la niche"; ws["B2"].font=Font(bold=True,size=18,color=GREEN)
    labels={"B3":"Mode","B4":"Marché","E3":"Statut","E4":"Date","B5":"Notes"}
    for k,v in labels.items(): ws[k]=v; ws[k].font=Font(bold=True,color=INK); ws[k].fill=fill(CREAM); ws[k].border=border
    for k in ["C3","C4","F3","F4"]: ws[k].fill=fill("FFFFFF"); ws[k].border=border
    ws.merge_cells("C5:K5"); ws["C5"].border=border
    ws["C3"]=example["mode"] if example else "PRODUIT PUR"; ws["C4"]=example["market"] if example else "FR"
    ws["F3"]=example["status"] if example else "Idée"; ws["F4"]=example["date"] if example else None
    ws["C5"]=example.get("notes","") if example else ""
    ws["H3"]="Seuils"; ws["H3"].font=Font(bold=True,color=GREY,size=9)
    ws["I3"]="PRODUIT PUR"; ws["J3"]=12500; ws["I4"]="UNIVERS"; ws["J4"]=37500
    for k in ["I3","I4"]: ws[k].font=Font(size=9,color=GREY)
    for k in ["J3","J4"]: ws[k].number_format='# ##0'; ws[k].font=Font(size=9,color=GREY)
    ws["K3"]="Bande REVIEW"; ws["K4"]=0.8; ws["K3"].font=Font(size=9,color=GREY); ws["K4"].number_format='0%'; ws["K4"].font=Font(size=9,color=GREY)
    hdr(ws,7,["Niveau","Arborescence","Mot-clé","Volume / mois","Seuil","Verdict","Lien AliExpress","Prix AliExpress","Prix cible","Marge brute","Notes / source"])
    ws.freeze_panes="A8"
    dvl=DataValidation(type="list",formula1='"Mot-clé principal,Collection,Produit"',allow_blank=True); ws.add_data_validation(dvl); dvl.add("A8:A80")
    dvm=DataValidation(type="list",formula1='"PRODUIT PUR,UNIVERS"',allow_blank=True); ws.add_data_validation(dvm); dvm.add("C3")
    dvk=DataValidation(type="list",formula1='"FR,UK,DE,ES,IT,US"',allow_blank=True); ws.add_data_validation(dvk); dvk.add("C4")
    dvs=DataValidation(type="list",formula1='"Idée,Mesure en cours,PASS,REVIEW,STOP,Sourcing,GO_FINAL,WATCH_FINAL,NO_GO_FINAL"',allow_blank=True); ws.add_data_validation(dvs); dvs.add("F3")
    for r in range(8,81):
        ws.cell(row=r,column=2,value=f'=IF(A{r}="","",IF(A{r}="Mot-clé principal","◆ "&C{r},IF(A{r}="Collection","    ├─ "&C{r},"         · "&C{r})))')
        ws.cell(row=r,column=5,value=f'=IF(A{r}="","",IF(A{r}="Mot-clé principal",IF($C$3="UNIVERS",$J$4,$J$3),IF(A{r}="Collection",$J$3,IF($C$3="PRODUIT PUR",$J$3,""))))')
        ws.cell(row=r,column=6,value=f'=IF(OR(D{r}="",E{r}=""),"",IF(D{r}>=E{r},"PASS",IF(D{r}>=E{r}*$K$4,"REVIEW","STOP")))')
        ws.cell(row=r,column=10,value=f'=IF(OR(H{r}="",I{r}=""),"",I{r}/1.2-H{r})')
        for c in range(1,12):
            cell=ws.cell(row=r,column=c); cell.border=border
        ws.cell(row=r,column=4).number_format='# ##0'; ws.cell(row=r,column=5).number_format='# ##0'
        for c in (8,9,10): ws.cell(row=r,column=c).number_format='#,##0.00" €"' if (example is None or example["market"]=="FR") else '"£"#,##0.00'
        ws.cell(row=r,column=7).font=Font(color="2563EB",underline="single"); ws.cell(row=r,column=2).font=Font(color=INK)
    # row styling by level
    ws.conditional_formatting.add("A8:K80",FormulaRule(formula=['$A8="Mot-clé principal"'],fill=fill(GREEN),font=Font(bold=True,color="FFFFFF")))
    ws.conditional_formatting.add("A8:K80",FormulaRule(formula=['$A8="Collection"'],fill=fill(PALE),font=Font(bold=True,color=INK)))
    ws.conditional_formatting.add("F8:F80",CellIsRule(operator="equal",formula=['"PASS"'],fill=fill(GRN),font=Font(bold=True,color="166534")))
    ws.conditional_formatting.add("F8:F80",CellIsRule(operator="equal",formula=['"REVIEW"'],fill=fill(AMB),font=Font(bold=True,color="92400E")))
    ws.conditional_formatting.add("F8:F80",CellIsRule(operator="equal",formula=['"STOP"'],fill=fill(RED),font=Font(bold=True,color="991B1B")))
    ws.conditional_formatting.add("J8:J80",CellIsRule(operator="lessThan",formula=['0'],font=Font(color="991B1B",bold=True)))
    if example:
        for i,row in enumerate(example["rows"]):
            r=8+i; lvl,kw,vol,link,pa,pc,note=row
            ws.cell(row=r,column=1,value=lvl); ws.cell(row=r,column=3,value=kw); ws.cell(row=r,column=4,value=vol)
            if link: ws.cell(row=r,column=7,value=f'=HYPERLINK("{link}","{link.replace("https://","")[:48]}")')
            if pa is not None: ws.cell(row=r,column=8,value=pa)
            if pc is not None: ws.cell(row=r,column=9,value=pc)
            ws.cell(row=r,column=11,value=note)
    return ws

niche_sheet("🧩 MODÈLE",GREY)
ex={"name":"Abri chat extérieur (UK)","mode":"PRODUIT PUR","market":"UK","status":"Sourcing","date":"2026-09-08",
    "notes":"Exemple rempli avec la recherche UK du 07/09 : volumes DataForSEO United Kingdom, une variante L grise qualifiée GB 6–10 j.",
    "rows":[
     ("Mot-clé principal","outdoor cat house",14800,None,None,None,"MAX du groupe (outdoor cat house / shelter). Concurrents Search : Home & Roost."),
     ("Collection","insulated cat house",6600,None,None,None,"Famille isolée, non chauffée"),
     ("Produit","insulated outdoor cat house, raised, grey L",None,"https://www.aliexpress.com/item/1005010759559289.html",31.0,89.0,"Stone's Store — seule variante qualifiée GB 6–10 j"),
     ("Produit","outdoor cat shelter white, arch door",None,"https://www.aliexpress.com/item/1005009960200806.html",29.0,79.0,"Backup, même magasin, EU warehouses refusés GB"),
     ("Collection","heated cat house outdoor",5400,None,None,None,"Version chauffée : non sourcée (câble, conformité UK)"),
     ("Produit","heated outdoor cat house 240V UK plug",None,"",None,None,"À sourcer seulement si le non chauffé passe"),
     ("Collection","cat house accessories",1300,None,None,None,"Coussin, rabat de porte, tapis chauffant"),
     ("Produit","self-warming cat pad",None,"",None,None,"Complémentaire possible"),
    ]}
niche_sheet("Abri chat extérieur (UK)",GREEN,ex)
ix["A5"]="Abri chat extérieur (UK)"; ix["M5"]="Exemple rempli (recherche UK 07/09)"
ix.add_data_validation(dv)
wb.move_sheet("🗂 Index",offset=-0)
out="/private/tmp/claude-502/-Users-Hakim-Documents-Boutiques-drop/f13db31b-0454-40da-9b2d-d09781f8f692/scratchpad/Idees-de-niches.xlsx"
wb.save(out); print("saved",out)
